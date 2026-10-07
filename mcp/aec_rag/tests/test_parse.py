"""Tests for the parsing paths (tasks 2.1, 2.4, 2.5, 3H.13, 3H.14).

The format tests run the real readers — a stub would only prove that our own
dataclasses work, and the thing worth pinning here is what pdfium, RapidOCR and
Docling actually do with page numbers and headings. The converter is
session-scoped because loading the layout model costs far more than the
conversions themselves.
"""

from __future__ import annotations

import docx
import openpyxl
import pytest
from aec_rag.parse import (
    ParseError,
    build_converter,
    parse_docx,
    parse_pdf,
    parse_spreadsheet,
    parse_whole,
    recognize_pages,
    split_by_headings,
    usable_text_layer,
)
from pdf_fixture import build_scanned_pdf, build_text_pdf

# Each page carries a marker that appears nowhere else, so "page 2's text ended
# up under page 3" is a test failure and not a judgement call. Page 3 is blank
# on purpose: a page with nothing on it must still hold its place in the
# sequence, or every citation after it points one page short.
PAGE_LINES = [
    ["Installed power report", "The installed power is 125 kW."],
    ["Cable list", "Cable A: 3x50 mm2"],
    [],
    ["Notes", "Voltage is 0.4 kV."],
]
PAGE_MARKERS = ["125 kW", "3x50 mm2", None, "0.4 kV"]


@pytest.fixture(scope="session")
def converter():
    return build_converter()


@pytest.fixture(scope="session")
def sample_pdf(tmp_path_factory):
    path = tmp_path_factory.mktemp("pdf") / "sample.pdf"
    path.write_bytes(build_text_pdf(PAGE_LINES))
    return path


@pytest.fixture(scope="session")
def parsed_pdf(sample_pdf):
    # The fixture's pages are a line or two long; the default threshold is set
    # for real pages (task 3H.11) and has its own tests below.
    return parse_pdf(sample_pdf, min_letters=5)


@pytest.fixture(scope="session")
def sample_docx(tmp_path_factory):
    """A Russian DOCX with two heading levels, a table and a preamble above the
    first heading — the shape a real design assignment actually has.
    """
    path = tmp_path_factory.mktemp("docx") / "задание.docx"
    document = docx.Document()
    document.add_paragraph("Редакция 2 от 03.02.2026.")
    document.add_heading("Задание ЭОМ", level=1)
    document.add_paragraph("Установленная мощность 125 кВт.")
    document.add_heading("Кабельный журнал", level=2)
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Позиция"
    table.cell(0, 1).text = "Значение"
    table.cell(1, 0).text = "Кабель А"
    table.cell(1, 1).text = "3х50 мм2"
    document.add_heading("Примечания", level=2)
    document.add_paragraph("Резерв 15%.")
    document.save(path)
    return path


@pytest.fixture(scope="session")
def parsed_docx(converter, sample_docx):
    return parse_docx(sample_docx, converter=converter)


def test_pdf_reports_every_page_with_its_1_based_number(parsed_pdf):
    assert parsed_pdf.doctype == "pdf"
    assert [segment.locator for segment in parsed_pdf.segments] == [
        {"page": 1},
        {"page": 2},
        {"page": 3},
        {"page": 4},
    ]


def test_pdf_marker_lands_on_its_own_page_and_nowhere_else(parsed_pdf):
    for segment, marker in zip(parsed_pdf.segments, PAGE_MARKERS, strict=True):
        if marker is None:
            continue
        assert marker in segment.markdown
        others = [other for other in parsed_pdf.segments if other is not segment]
        assert all(marker not in other.markdown for other in others)


def test_pdf_blank_page_keeps_its_place_in_the_sequence(parsed_pdf):
    assert parsed_pdf.segments[2].markdown == ""


def test_pdf_page_without_a_layer_waits_for_recognition(parsed_pdf):
    assert parsed_pdf.pending_pages == (3,)


def test_pdf_document_markdown_holds_the_whole_document(parsed_pdf):
    for marker in PAGE_MARKERS:
        if marker is not None:
            assert marker in parsed_pdf.markdown


# What a frame stamp on a scanned sheet reads like: measured 2026-10-07, such a
# page carries 64–69 letters and digits (PLAN §3H item C).
STAMP = "Взам. инв. № Подп. и дата Инв. № подл. Изм. Кол.уч. Лист №док. Подп. Дата 02-ПЗ"
PROSE = (
    "Инженерно-геологические условия площадки определены по результатам бурения "
    "двенадцати скважин глубиной до тридцати метров. Грунты основания представлены "
    "техногенными отложениями, песками средней крупности и суглинками тугопластичными. "
)


def test_text_layer_of_real_prose_is_usable():
    assert usable_text_layer(PROSE)


def test_a_frame_stamp_alone_is_not_a_text_layer():
    assert not usable_text_layer(STAMP)


def test_a_broken_font_encoding_is_not_a_text_layer():
    """Plenty of characters, few letters — what a font without a ToUnicode map
    extracts as."""
    assert not usable_text_layer("".join("§¤¦¬" for _ in range(200)) + "аб12")


def test_pdf_with_a_full_page_of_text_needs_no_recognition(tmp_path):
    line = "The installed power of the building is 125 kW at 0.4 kV."
    path = tmp_path / "text.pdf"
    path.write_bytes(build_text_pdf([[line] * 8, ["Sheet 2"]]))

    parsed = parse_pdf(path)

    assert "125 kW" in parsed.segments[0].markdown
    assert parsed.segments[1].markdown == ""
    assert parsed.pending_pages == (2,)


SCAN_LINES = [
    ["Installed power", "125 kW total"],
    ["Cable list", "Cable A 3x50"],
]


@pytest.fixture(scope="session")
def scan_pdf(tmp_path_factory):
    path = tmp_path_factory.mktemp("scan") / "scan.pdf"
    path.write_bytes(build_scanned_pdf(SCAN_LINES))
    return path


def test_scanned_pdf_has_no_text_layer_to_begin_with(scan_pdf):
    """Guards the fixture itself: if it ever grew a text layer the OCR test
    below would pass without OCR running at all.
    """
    parsed = parse_pdf(scan_pdf, min_letters=1)

    assert [segment.markdown for segment in parsed.segments] == ["", ""]
    assert parsed.pending_pages == (1, 2)


def test_scanned_pdf_is_recognised_page_by_page(scan_pdf):
    results = dict(recognize_pages(scan_pdf, [1, 2]))

    assert "125" in results[1]
    assert "125" not in results[2]
    assert "3x50" in results[2].replace(" ", "")


def test_a_page_that_cannot_be_recognised_does_not_end_the_file(scan_pdf):
    results = list(recognize_pages(scan_pdf, [1, 9, 2]))

    assert [page for page, _ in results] == [1, 9, 2]
    assert isinstance(results[1][1], ParseError)
    assert "125" in results[0][1]
    assert isinstance(results[2][1], str)


def test_a_caller_that_needs_the_text_now_gets_both_passes(scan_pdf):
    """`wiki_add_raw` has no background run to finish the scanned pages later."""
    parsed = parse_whole(scan_pdf)

    assert "125" in parsed.segments[0].markdown
    assert parsed.pending_pages == ()
    assert "125" in parsed.markdown


def test_pdf_under_a_cyrillic_path_is_parsed(tmp_path):
    """Our users' files are named in Russian and live in Russian-named folders.
    Docling's backend could not open such a path at all (measured 2026-08-14);
    the text layer is read from bytes, so the path never reaches pdfium.
    """
    directory = tmp_path / "проекты" / "ЖК Нагатинский"
    directory.mkdir(parents=True)
    path = directory / "02-ПЗ-Нагатинская 297-1.pdf"
    path.write_bytes(build_text_pdf([["Installed power report", "125 kW"]]))

    parsed = parse_pdf(path, min_letters=1)

    assert "125 kW" in parsed.markdown


def test_unreadable_file_raises_parse_error(tmp_path):
    broken = tmp_path / "broken.pdf"
    broken.write_bytes(b"%PDF-1.4\nthis is not a pdf\n")

    with pytest.raises(ParseError) as excinfo:
        parse_pdf(broken)

    assert str(broken) in str(excinfo.value)


def test_docx_segments_are_headings_not_pages(parsed_docx):
    assert parsed_docx.doctype == "docx"
    assert [segment.locator for segment in parsed_docx.segments] == [
        {},
        {"heading": "Задание ЭОМ"},
        {"heading": "Кабельный журнал"},
        {"heading": "Примечания"},
    ]


def test_docx_preamble_above_the_first_heading_survives(parsed_docx):
    assert "Редакция 2" in parsed_docx.segments[0].markdown


def test_docx_text_stays_under_its_own_heading(parsed_docx):
    by_heading = {
        segment.locator.get("heading"): segment.markdown for segment in parsed_docx.segments
    }
    assert "125 кВт" in by_heading["Задание ЭОМ"]
    assert "3х50 мм2" in by_heading["Кабельный журнал"]
    assert "Резерв 15%" in by_heading["Примечания"]


def test_docx_table_is_extracted_as_a_markdown_table(parsed_docx):
    cable = next(
        segment
        for segment in parsed_docx.segments
        if segment.locator.get("heading") == "Кабельный журнал"
    )
    assert "| Позиция" in cable.markdown
    assert "3х50 мм2" in cable.markdown


def test_docx_segment_repeats_its_own_heading(parsed_docx):
    for segment in parsed_docx.segments:
        heading = segment.locator.get("heading")
        if heading is not None:
            assert heading in segment.markdown


@pytest.fixture(scope="session")
def sample_xlsx(tmp_path_factory):
    """Three sheets: one ordinary, one empty, one whose table does not start at
    A1 — the range is only worth anything if it points at where the data is.
    """
    path = tmp_path_factory.mktemp("xlsx") / "ведомость.xlsx"
    book = openpyxl.Workbook()
    first = book.active
    first.title = "Ведомость"
    first.append(["Позиция", "Мощность, кВт"])
    first.append(["Насос К-1", 55])
    first.append(["Насос К-2", 70])
    book.create_sheet("Пустой")
    third = book.create_sheet("Кабели")
    third["C5"] = "Марка"
    third["D5"] = "Сечение"
    third["C6"] = "ВВГнг"
    third["D6"] = "3х50"
    book.save(path)
    return path


@pytest.fixture(scope="session")
def parsed_xlsx(converter, sample_xlsx):
    return parse_spreadsheet(sample_xlsx, converter=converter)


def test_xlsx_segment_names_its_sheet_and_range(parsed_xlsx):
    assert parsed_xlsx.doctype == "xlsx"
    assert [segment.locator for segment in parsed_xlsx.segments] == [
        {"sheet": "Ведомость", "range": "A1:B3"},
        {"sheet": "Кабели", "range": "C5:D6"},
    ]


def test_xlsx_empty_sheet_produces_no_segment(parsed_xlsx):
    assert "Пустой" not in [segment.locator.get("sheet") for segment in parsed_xlsx.segments]


def test_xlsx_table_content_stays_with_its_sheet(parsed_xlsx):
    by_sheet = {segment.locator["sheet"]: segment.markdown for segment in parsed_xlsx.segments}
    assert "Насос К-1" in by_sheet["Ведомость"]
    assert "Насос К-1" not in by_sheet["Кабели"]
    assert "ВВГнг" in by_sheet["Кабели"]


def test_csv_is_one_table_the_citation_can_only_name_the_file(converter, tmp_path):
    path = tmp_path / "нагрузки.csv"
    path.write_text("Позиция,Мощность\nНасос К-1,55\n", encoding="utf-8")

    parsed = parse_spreadsheet(path, converter=converter)

    assert parsed.doctype == "csv"
    assert [segment.locator for segment in parsed.segments] == [{}]
    assert "Насос К-1" in parsed.segments[0].markdown


def test_split_by_headings_returns_one_section_when_there_are_none():
    assert split_by_headings("просто текст\nбез заголовков") == [
        (None, "просто текст\nбез заголовков")
    ]


def test_split_by_headings_ignores_hashes_inside_a_fenced_block():
    markdown = "# Раздел\n\n```\n# это комментарий\n```\n\nхвост"

    sections = split_by_headings(markdown)

    assert [heading for heading, _ in sections] == ["Раздел"]
    assert "# это комментарий" in sections[0][1]


def test_split_by_headings_drops_an_empty_preamble():
    assert split_by_headings("\n\n## Первый\n\nтекст") == [("Первый", "## Первый\n\nтекст")]
