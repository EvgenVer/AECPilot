"""Собирает документы мини-проекта из их текстовых исходников (задача 4.14).

    uv run python eval/fixtures/wiki-project/build_docs.py

Исходники лежат рядом в `src/` обычным markdown — так их читает и правит
человек. Индексатор берёт только `.pdf`, `.docx`, `.xlsx`, `.csv`
(`aec_rag.parse.SUPPORTED_SUFFIXES`), поэтому в `.docx` они превращаются здесь,
кодом, а не руками: собранный документ обязан слово в слово совпадать со своим
исходником, иначе эвал меряет не то, что написано.
"""

from __future__ import annotations

from pathlib import Path

from docx import Document

HERE = Path(__file__).resolve().parent
SRC = HERE / "src"


def build(source: Path) -> Path:
    document = Document()
    for line in source.read_text(encoding="utf-8").split("\n"):
        stripped = line.strip()
        if not stripped:
            continue
        if stripped.startswith("## "):
            document.add_heading(stripped[3:], level=2)
        elif stripped.startswith("# "):
            document.add_heading(stripped[2:], level=1)
        else:
            document.add_paragraph(stripped)

    target = HERE / f"{source.stem}.docx"
    document.save(target)
    return target


def main() -> int:
    for source in sorted(SRC.glob("*.md")):
        print(build(source).name)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
